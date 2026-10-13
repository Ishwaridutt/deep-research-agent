from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from src.models.document_models import (
    DocumentFormat,
    DocumentSourceLocation,
    ParsedDocument,
    ParsedDocumentBlock,
)


class DocumentParser:
    """Extract normalized text blocks from supported document formats."""

    def parse(
        self,
        document_id: str,
        path: Path,
        document_format: DocumentFormat,
    ) -> ParsedDocument:
        """Parse one stored document while retaining available source locations."""

        if document_format == DocumentFormat.PDF:
            blocks = [
                ParsedDocumentBlock(
                    text=text.strip(),
                    location=DocumentSourceLocation(
                        page_start=i,
                        page_end=i,
                    ),
                )
                for i, page in enumerate(PdfReader(path).pages, 1)
                if (text := page.extract_text())
            ]

        elif document_format == DocumentFormat.DOCX:
            blocks = [
                ParsedDocumentBlock(text=p.text.strip())
                for p in DocxDocument(path).paragraphs
                if p.text.strip()
            ]

        else:
            text = path.read_text(
                encoding="utf-8",
                errors="replace",
            ).strip()

            blocks = (
                [ParsedDocumentBlock(text=text)]
                if text
                else []
            )

        if not blocks:
            raise ValueError(
                "No extractable text found in document."
            )

        return ParsedDocument(
            document_id=document_id,
            blocks=blocks,
        )
    
